"""Dekodowanie pliku źródłowego do czystego, ponumerowanego tekstu + metadanych.

Formaty obsługiwane bez dodatkowych zależności: `.md` / `.txt` z frontmatterem
(from/to/subject/date/message_id/in_reply_to) — to jest już „zdekodowany” kształt, taki sam,
jaki produkuje dekodowanie `.msg` niżej. Pliki demo w `_input/originals/` mają ten format
wprost, żeby testować resztę przepływu bez zależności od Outlooka.

`.msg`, `.docx` i `.pptx` wymagają odpowiednio `extract_msg`, `python-docx` i `python-pptx` —
importowane dopiero przy użyciu, nie są twardą zależnością tego workspace'u.
"""
from __future__ import annotations
from dataclasses import dataclass, field
from pathlib import Path
import re

_FRONTMATTER_RE = re.compile(r"^---\n(.*?)\n---\n?(.*)$", re.DOTALL)


@dataclass
class DecodedDocument:
    metadata: dict
    text: str

    def numbered_text(self) -> str:
        lines = self.text.splitlines()
        return "\n".join(f"{i + 1:>4}| {line}" for i, line in enumerate(lines))

    def evidence(self, start: int, end: int) -> str:
        """1-indeksowane, włącznie z końcem — dokładnie te linie, które wskazał model."""
        lines = self.text.splitlines()
        return "\n".join(lines[start - 1:end])


def _parse_frontmatter(raw: str) -> tuple[dict, str]:
    match = _FRONTMATTER_RE.match(raw)
    if not match:
        return {}, raw
    header, body = match.groups()
    metadata = {}
    for line in header.splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            metadata[k.strip()] = v.strip()
    return metadata, body.lstrip("\n")


def decode(path: Path) -> DecodedDocument:
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt"):
        metadata, body = _parse_frontmatter(path.read_text(encoding="utf-8"))
        return DecodedDocument(metadata=metadata, text=body)
    if suffix == ".msg":
        return _decode_msg(path)
    if suffix == ".docx":
        return _decode_docx(path)
    if suffix == ".pptx":
        return _decode_pptx(path)
    raise ValueError(f"Nieobsługiwany format źródła: {suffix}")


def _decode_msg(path: Path) -> DecodedDocument:
    try:
        import extract_msg
    except ImportError as exc:
        raise RuntimeError(
            "Dekodowanie .msg wymaga biblioteki extract_msg (pip install extract-msg)."
        ) from exc
    from bs4 import BeautifulSoup

    msg = extract_msg.Message(str(path))
    body = msg.body or ""
    if not body and msg.htmlBody:
        body = BeautifulSoup(msg.htmlBody, "html.parser").get_text("\n")
    metadata = {
        "from": msg.sender or "—",
        "to": msg.to or "—",
        "subject": msg.subject or "—",
        "date": str(msg.date) if msg.date else "—",
        "message_id": msg.messageId or "—",
        "in_reply_to": getattr(msg, "inReplyTo", "") or "—",
    }
    return DecodedDocument(metadata=metadata, text=body.strip())


def _decode_docx(path: Path) -> DecodedDocument:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError(
            "Dekodowanie .docx wymaga biblioteki python-docx (pip install python-docx)."
        ) from exc
    document = docx.Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs)
    return DecodedDocument(metadata=_office_metadata(document.core_properties), text=text)


def _decode_pptx(path: Path) -> DecodedDocument:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "Dekodowanie .pptx wymaga biblioteki python-pptx (pip install python-pptx)."
        ) from exc

    presentation = Presentation(str(path))
    lines: list[str] = []
    for i, slide in enumerate(presentation.slides, start=1):
        lines.append(f"[Slajd {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs)
                    if text.strip():
                        lines.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    if row_text.strip():
                        lines.append(row_text)
    return DecodedDocument(metadata=_office_metadata(presentation.core_properties), text="\n".join(lines))


_TOOL_GENERATED_AUTHORS = {
    "python-docx", "pptxgenjs", "un-named", "unknown", "microsoft office user",
}


def _office_metadata(core_properties) -> dict:
    """Data z metadanych pliku .docx/.pptx — brak nagłówka e-maila, więc bierzemy datę
    modyfikacji dokumentu (ostatniego zapisu) jako przybliżenie daty zdarzenia; etap `read`
    (model) może ją nadpisać, gdy tekst wprost mówi o innej dacie.

    `author` bywa nazwą biblioteki, która zapisała plik (np. "python-docx", "PptxGenJS"), nie
    prawdziwym autorem — takie wartości odrzucamy, żeby nie pokazywać ich jako "od kogo".
    """
    date = core_properties.modified or core_properties.created
    metadata = {"date": str(date) if date else "—"}
    author = (core_properties.author or "").strip()
    if author and author.lower() not in _TOOL_GENERATED_AUTHORS:
        metadata["from"] = author
    if core_properties.subject:
        metadata["subject"] = core_properties.subject
    return metadata
